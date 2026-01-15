#!/usr/bin/env python3

import argparse
import csv
import json
import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen

REQUIRED_COLUMNS = {
    "id",
    "name",
    "halal_likelihood",
    "halal_type",
    "halal_confidence",
    "halal_reasoning",
}

VALIDATION_DIR = Path(__file__).resolve().parent.parent / "data"
REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"

CERTIFIER_PATTERNS = {
    "SBNY": ["sbny", "shariah board", "shariah board ny"],
    "HMS": ["hms"],
    "HFSAA": ["hfsaa"],
    "IFANCA": ["ifanca", "islamic food and nutrition council"],
}

OFFICIAL_KEYWORDS = [
    "official",
    "website",
    "menu",
    "sign",
    "signage",
    "owner",
    "staff",
    "phone",
    "call",
    "email",
    "instagram",
    "facebook",
]

DIRECTORY_KEYWORDS = [
    "halaltrip",
    "muslim pro",
    "muslimpro",
    "halal guide",
    "halalfood",
    "crescent",
    "halal directory",
]

REVIEW_KEYWORDS = [
    "review",
    "reviews",
    "yelp",
    "google review",
    "tripadvisor",
]

NEGATIVE_KEYWORDS = [
    "not halal",
    "non-halal",
    "non halal",
]

PORK_KEYWORDS = [
    "pork",
    "bacon",
    "ham",
    "pepperoni",
]

ALCOHOL_KEYWORDS = [
    "serves alcohol",
    "alcohol",
    "beer",
    "wine",
    "cocktail",
    "liquor",
]

URL_PATTERN = re.compile(r"https?://\S+|www\.\S+", re.IGNORECASE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Apply halal validation results to Supabase.")
    parser.add_argument("--file", required=True, help="CSV/XLSX filename in data/ or a path")
    parser.add_argument("--db-url", dest="db_url", help="Postgres connection string")
    parser.add_argument("--use-rest", action="store_true", help="Use Supabase REST API")
    parser.add_argument("--supabase-url", dest="supabase_url", help="Supabase URL")
    parser.add_argument("--supabase-key", dest="supabase_key", help="Supabase service role key")
    parser.add_argument("--apply", action="store_true", help="Apply updates to the database")
    return parser.parse_args()


def resolve_file_path(file_arg: str) -> Path:
    candidate = Path(file_arg)
    if candidate.is_absolute() or candidate.exists():
        return candidate
    return VALIDATION_DIR / file_arg


def normalize_header(header: str) -> str:
    return header.strip().lower()


def load_csv_rows(file_path: Path) -> List[Dict[str, object]]:
    rows = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            raw_headers = next(reader)
        except StopIteration:
            return rows
        headers = [normalize_header(str(h)) for h in raw_headers]
        for row in reader:
            if not any(str(cell).strip() for cell in row if cell is not None):
                continue
            row_dict = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
            rows.append(row_dict)
    return rows


def load_xlsx_rows(file_path: Path) -> List[Dict[str, object]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for .xlsx files. Install with: pip install openpyxl"
        ) from exc

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return rows
    headers = [normalize_header(str(h or "")) for h in raw_headers]
    for row in iterator:
        if not any(str(cell).strip() for cell in row if cell is not None):
            continue
        row_dict = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        rows.append(row_dict)
    return rows


def load_rows(file_path: Path) -> List[Dict[str, object]]:
    if file_path.suffix.lower() == ".csv":
        return load_csv_rows(file_path)
    if file_path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_xlsx_rows(file_path)
    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def get_db_url(cli_value: Optional[str]) -> str:
    if cli_value:
        return cli_value
    for key in ("DATABASE_URL", "SUPABASE_DB_URL", "SUPABASE_DATABASE_URL"):
        value = os.getenv(key)
        if value:
            return value
    raise RuntimeError("Database URL not found. Use --db-url or set DATABASE_URL.")


def get_db_module():
    try:
        import psycopg2

        return "psycopg2", psycopg2
    except ImportError:
        pass

    try:
        import psycopg

        return "psycopg", psycopg
    except ImportError as exc:
        raise RuntimeError(
            "psycopg2 or psycopg is required. Install with: pip install psycopg2-binary"
        ) from exc


def detect_table_name(cursor) -> str:
    cursor.execute(
        """
        select table_name
        from information_schema.tables
        where table_schema = 'public'
          and table_name in ('place', 'places')
        """
    )
    names = {row[0] for row in cursor.fetchall()}
    if "place" in names:
        return "place"
    if "places" in names:
        return "places"
    raise RuntimeError("Neither public.place nor public.places exists in the database.")


def get_supabase_credentials(args: argparse.Namespace) -> Tuple[str, str]:
    url = args.supabase_url or os.getenv("SUPABASE_URL")
    key = args.supabase_key or os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        raise RuntimeError("Supabase URL/key missing. Use --supabase-url/--supabase-key.")
    return url.rstrip("/"), key


def rest_request(
    method: str, url: str, headers: Dict[str, str], payload: Optional[Dict[str, object]] = None
) -> Tuple[int, str]:
    data = None
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
    request = Request(url, data=data, headers=headers, method=method)
    try:
        with urlopen(request) as response:
            body = response.read().decode("utf-8")
            return response.status, body
    except HTTPError as exc:
        body = exc.read().decode("utf-8")
        return exc.code, body


def detect_table_name_rest(base_url: str, headers: Dict[str, str]) -> str:
    for table_name in ("place", "places"):
        status, _ = rest_request(
            "GET", f"{base_url}/rest/v1/{table_name}?select=id&limit=1", headers
        )
        if status in (200, 206):
            return table_name
        if status in (401, 403):
            raise RuntimeError("Supabase REST auth failed.")
    raise RuntimeError("Neither public.place nor public.places exists in Supabase REST.")


def rest_update_row(
    base_url: str,
    headers: Dict[str, str],
    table_name: str,
    row_id: str,
    payload: Dict[str, object],
) -> List[Dict[str, object]]:
    encoded_id = quote(row_id, safe="")
    url = f"{base_url}/rest/v1/{table_name}?id=eq.{encoded_id}&select=name,halal_status"
    update_headers = dict(headers)
    update_headers["Prefer"] = "return=representation"
    status, body = rest_request("PATCH", url, update_headers, payload)
    if status in (200, 201):
        return json.loads(body) if body else []
    if status == 204:
        return []
    raise RuntimeError(f"REST update failed with status {status}: {body}")


def parse_int(value: object) -> Optional[int]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    cleaned = text.replace("%", "")
    try:
        return int(float(cleaned))
    except ValueError:
        return None


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_is_zabiha(reasoning: str) -> Optional[bool]:
    lowered = reasoning.lower()
    if "zabiha" not in lowered and "zabihah" not in lowered:
        return None
    negative_patterns = [
        "not listed on zabiha",
        "not on zabiha",
        "no zabiha",
        "not listed in zabiha",
        "not listed on zabihah",
    ]
    if any(pattern in lowered for pattern in negative_patterns):
        return False
    return True


def extract_certifier_org(reasoning: str) -> Optional[str]:
    lowered = reasoning.lower()
    certified_segment = None
    match = re.search(r"certified by[^.]*", lowered)
    if match:
        certified_segment = match.group(0)
    if certified_segment:
        for org, patterns in CERTIFIER_PATTERNS.items():
            if any(pattern in certified_segment for pattern in patterns):
                return org
    earliest = None
    earliest_org = None
    for org, patterns in CERTIFIER_PATTERNS.items():
        for pattern in patterns:
            idx = lowered.find(pattern)
            if idx != -1 and (earliest is None or idx < earliest):
                earliest = idx
                earliest_org = org
    return earliest_org


def strip_urls(text: str) -> str:
    return URL_PATTERN.sub("", text).strip()


def pick_negative_evidence(lowered: str) -> Optional[str]:
    if any(pattern in lowered for pattern in NEGATIVE_KEYWORDS):
        return "Marked as not halal."
    if any(pattern in lowered for pattern in PORK_KEYWORDS):
        return "Menu mentions pork."
    if any(pattern in lowered for pattern in ALCOHOL_KEYWORDS):
        return "Serves alcohol."
    return None


def select_evidence(
    halal_likelihood_norm: str, reasoning: str, certifier_org: Optional[str]
) -> Optional[str]:
    lowered = reasoning.lower()

    if halal_likelihood_norm == "LIKELY_NOT_HALAL":
        negative = pick_negative_evidence(lowered)
        return negative or "Evidence suggests not halal."

    if certifier_org:
        return f"Certified by {certifier_org}."
    if "certified" in lowered or "certification" in lowered:
        return "Certified halal."
    if any(keyword in lowered for keyword in OFFICIAL_KEYWORDS):
        return "Official info indicates halal."
    if "zabiha" in lowered or "zabihah" in lowered:
        return "Listed on Zabiha."
    if any(keyword in lowered for keyword in DIRECTORY_KEYWORDS):
        return "Listed in a halal directory."
    if any(keyword in lowered for keyword in REVIEW_KEYWORDS):
        return "Reviews mention halal."
    return None


def label_for_note(cc_halal_status: str) -> str:
    if cc_halal_status == "only":
        return "Fully"
    if cc_halal_status == "yes":
        return "Options"
    return "Unclear"


def truncate(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def build_cc_note(
    cc_halal_status: str,
    halal_likelihood_norm: str,
    reasoning: str,
    certifier_org: Optional[str],
    is_zabiha: Optional[bool],
) -> str:
    label = label_for_note(cc_halal_status)
    evidence = select_evidence(halal_likelihood_norm, reasoning, certifier_org)
    sentences = [f"Halal: {label}."]

    if evidence:
        cleaned = strip_urls(evidence)
        if cleaned:
            sentences.append(cleaned if cleaned.endswith(".") else f"{cleaned}.")

    if is_zabiha is True and (not evidence or "zabiha" not in evidence.lower()):
        sentences.append("Listed on Zabiha.")

    if cc_halal_status == "unclear":
        sentences.append("No official menu/certification found.")

    note = " ".join(sentences)
    return truncate(note, 350)


def map_cc_status(halal_likelihood_norm: str, halal_type_norm: str) -> str:
    if halal_likelihood_norm == "LIKELY_HALAL" and halal_type_norm == "FULLY_HALAL":
        return "only"
    if halal_likelihood_norm == "LIKELY_HALAL" and halal_type_norm == "HALAL_OPTIONS_ONLY":
        return "yes"
    if halal_likelihood_norm == "LIKELY_NOT_HALAL":
        return "no"
    return "unclear"


def dedupe_rows(
    rows: List[Dict[str, object]]
) -> Tuple[List[Dict[str, object]], Dict[str, List[Dict[str, object]]]]:
    best_by_id: Dict[str, Dict[str, object]] = {}
    duplicates: Dict[str, List[Dict[str, object]]] = {}
    first_seen_index: Dict[str, int] = {}
    for row in rows:
        row_id = row["id"]
        if row_id not in first_seen_index:
            first_seen_index[row_id] = row["row_index"]
        current = best_by_id.get(row_id)
        if current is None:
            best_by_id[row_id] = row
            continue
        duplicates.setdefault(row_id, [current]).append(row)
        if row["halal_confidence"] > current["halal_confidence"]:
            best_by_id[row_id] = row
    unique_rows = sorted(best_by_id.values(), key=lambda r: first_seen_index[r["id"]])
    return unique_rows, duplicates


def write_applied_report(report_path: Path, rows: List[Dict[str, object]]) -> None:
    headers = [
        "id",
        "name_file",
        "name_db",
        "existing_halal_status",
        "new_cc_halal_status",
        "differs_from_existing",
        "cc_halal_likelihood",
        "cc_halal_type",
        "cc_halal_confidence",
        "cc_is_zabiha",
        "cc_certifier_org",
        "cc_note",
    ]
    with report_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(col, "") for col in headers])


def write_duplicates_report(
    report_path: Path, duplicates: Dict[str, List[Dict[str, object]]]
) -> None:
    headers = ["id", "row_index", "halal_confidence", "name_file", "kept"]
    rows = []
    for row_id, entries in duplicates.items():
        best = max(entries, key=lambda entry: entry["halal_confidence"])
        for entry in sorted(entries, key=lambda entry: entry["row_index"]):
            rows.append(
                {
                    "id": row_id,
                    "row_index": entry["row_index"],
                    "halal_confidence": entry["halal_confidence"],
                    "name_file": entry["name"],
                    "kept": "true" if entry is best else "false",
                }
            )

    with report_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(col, "") for col in headers])


def main() -> int:
    args = parse_args()
    if not args.apply:
        print("This script only runs with --apply.")
        return 2

    file_path = resolve_file_path(args.file)
    if not file_path.exists():
        print(f"File not found: {file_path}")
        return 1

    raw_rows = load_rows(file_path)
    if not raw_rows:
        print("No rows found in file.")
        return 1

    header_set = set(raw_rows[0].keys())
    missing_columns = REQUIRED_COLUMNS - header_set
    if missing_columns:
        print(f"Missing required columns: {', '.join(sorted(missing_columns))}")
        return 1

    parsed_rows = []
    invalid_rows = []
    for idx, row in enumerate(raw_rows, start=2):
        row_id = normalize_text(row.get("id"))
        name = normalize_text(row.get("name"))
        halal_likelihood = normalize_text(row.get("halal_likelihood"))
        halal_type = normalize_text(row.get("halal_type"))
        halal_confidence = parse_int(row.get("halal_confidence"))
        halal_reasoning_raw = row.get("halal_reasoning")
        halal_reasoning = "" if halal_reasoning_raw is None else str(halal_reasoning_raw)

        if not (row_id and name and halal_likelihood and halal_type and halal_reasoning):
            invalid_rows.append({"row_index": idx, "id": row_id, "name": name})
            continue
        if halal_confidence is None:
            invalid_rows.append({"row_index": idx, "id": row_id, "name": name})
            continue

        parsed_rows.append(
            {
                "row_index": idx,
                "id": row_id,
                "name": name,
                "halal_likelihood": halal_likelihood,
                "halal_type": halal_type,
                "halal_confidence": halal_confidence,
                "halal_reasoning_raw": halal_reasoning,
            }
        )

    if not parsed_rows:
        print("No valid rows to process after validation.")
        return 1

    deduped_rows, duplicates = dedupe_rows(parsed_rows)

    applied_rows = []
    missing_ids = []
    updated_ids = []

    if args.use_rest:
        base_url, api_key = get_supabase_credentials(args)
        rest_headers = {
            "apikey": api_key,
            "Authorization": f"Bearer {api_key}",
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        table_name = detect_table_name_rest(base_url, rest_headers)

        for row in deduped_rows:
            likelihood_norm = row["halal_likelihood"].strip().upper()
            type_norm = row["halal_type"].strip().upper()
            cc_halal_status = map_cc_status(likelihood_norm, type_norm)
            certifier_org = extract_certifier_org(row["halal_reasoning_raw"])
            is_zabiha = extract_is_zabiha(row["halal_reasoning_raw"])
            cc_note = build_cc_note(
                cc_halal_status,
                likelihood_norm,
                row["halal_reasoning_raw"],
                certifier_org,
                is_zabiha,
            )

            payload = {
                "cc_halal_status": cc_halal_status,
                "cc_halal_likelihood": row["halal_likelihood"],
                "cc_halal_type": row["halal_type"],
                "cc_halal_confidence": row["halal_confidence"],
                "cc_note": cc_note,
                "cc_reasoning_raw": row["halal_reasoning_raw"],
                "cc_is_zabiha": is_zabiha,
                "cc_certifier_org": certifier_org,
            }

            response_rows = rest_update_row(
                base_url, rest_headers, table_name, row["id"], payload
            )
            if not response_rows:
                missing_ids.append({"id": row["id"], "name": row["name"]})
                continue

            result = response_rows[0]
            name_db = result.get("name")
            existing_halal_status = result.get("halal_status")
            differs = str(existing_halal_status) != cc_halal_status
            applied_rows.append(
                {
                    "id": row["id"],
                    "name_file": row["name"],
                    "name_db": name_db,
                    "existing_halal_status": existing_halal_status,
                    "new_cc_halal_status": cc_halal_status,
                    "differs_from_existing": str(differs).lower(),
                    "cc_halal_likelihood": row["halal_likelihood"],
                    "cc_halal_type": row["halal_type"],
                    "cc_halal_confidence": row["halal_confidence"],
                    "cc_is_zabiha": "" if is_zabiha is None else str(is_zabiha).lower(),
                    "cc_certifier_org": certifier_org or "",
                    "cc_note": cc_note,
                }
            )
            updated_ids.append(row["id"])
    else:
        db_url = get_db_url(args.db_url)
        _, db_module = get_db_module()
        connection = db_module.connect(db_url)
        connection.autocommit = False
        cursor = connection.cursor()

        try:
            table_name = detect_table_name(cursor)
            update_sql = f"""
                update public.{table_name}
                set
                    cc_halal_status = %s,
                    cc_halal_likelihood = %s,
                    cc_halal_type = %s,
                    cc_halal_confidence = %s,
                    cc_note = %s,
                    cc_reasoning_raw = %s,
                    cc_is_zabiha = %s,
                    cc_certifier_org = %s
                where id = %s
                returning name, halal_status
            """

            for row in deduped_rows:
                likelihood_norm = row["halal_likelihood"].strip().upper()
                type_norm = row["halal_type"].strip().upper()
                cc_halal_status = map_cc_status(likelihood_norm, type_norm)
                certifier_org = extract_certifier_org(row["halal_reasoning_raw"])
                is_zabiha = extract_is_zabiha(row["halal_reasoning_raw"])
                cc_note = build_cc_note(
                    cc_halal_status,
                    likelihood_norm,
                    row["halal_reasoning_raw"],
                    certifier_org,
                    is_zabiha,
                )

                cursor.execute(
                    update_sql,
                    (
                        cc_halal_status,
                        row["halal_likelihood"],
                        row["halal_type"],
                        row["halal_confidence"],
                        cc_note,
                        row["halal_reasoning_raw"],
                        is_zabiha,
                        certifier_org,
                        row["id"],
                    ),
                )
                result = cursor.fetchone()
                if not result:
                    missing_ids.append({"id": row["id"], "name": row["name"]})
                    continue

                name_db, existing_halal_status = result
                differs = str(existing_halal_status) != cc_halal_status
                applied_rows.append(
                    {
                        "id": row["id"],
                        "name_file": row["name"],
                        "name_db": name_db,
                        "existing_halal_status": existing_halal_status,
                        "new_cc_halal_status": cc_halal_status,
                        "differs_from_existing": str(differs).lower(),
                        "cc_halal_likelihood": row["halal_likelihood"],
                        "cc_halal_type": row["halal_type"],
                        "cc_halal_confidence": row["halal_confidence"],
                        "cc_is_zabiha": "" if is_zabiha is None else str(is_zabiha).lower(),
                        "cc_certifier_org": certifier_org or "",
                        "cc_note": cc_note,
                    }
                )
                updated_ids.append(row["id"])

            expected_found = len(deduped_rows) - len(missing_ids)
            if len(updated_ids) != expected_found:
                connection.rollback()
                missing_update_ids = [
                    row["id"] for row in deduped_rows if row["id"] not in updated_ids
                ]
                print("Update count mismatch; rolling back.")
                print("Missing update IDs:")
                for missing_id in missing_update_ids:
                    print(f"- {missing_id}")
                return 1

            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            cursor.close()
            connection.close()

    expected_found = len(deduped_rows) - len(missing_ids)
    if len(updated_ids) != expected_found:
        missing_update_ids = [row["id"] for row in deduped_rows if row["id"] not in updated_ids]
        print("Update count mismatch; aborting.")
        print("Missing update IDs:")
        for missing_id in missing_update_ids:
            print(f"- {missing_id}")
        return 1

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    base_name = file_path.stem
    applied_report_path = REPORTS_DIR / f"{base_name}__applied.csv"
    write_applied_report(applied_report_path, applied_rows)

    if duplicates:
        duplicates_report_path = REPORTS_DIR / f"{base_name}__duplicates.csv"
        write_duplicates_report(duplicates_report_path, duplicates)

    unclear_count = sum(1 for row in applied_rows if row["new_cc_halal_status"] == "unclear")
    differs_count = sum(1 for row in applied_rows if row["differs_from_existing"] == "true")
    duplicate_count = sum(len(items) - 1 for items in duplicates.values())

    print("Applied report:", applied_report_path)
    print("File rows total:", len(raw_rows))
    print("Rows after validation:", len(parsed_rows))
    print("Rows after dedupe:", len(deduped_rows))
    print("Updated rows count:", len(applied_rows))
    print("Missing ids count:", len(missing_ids))
    print("Duplicates count:", duplicate_count)
    print("Unclear status count:", unclear_count)
    print("Differs from existing count:", differs_count)

    if invalid_rows:
        print("Invalid rows skipped:")
        for row in invalid_rows:
            print(f"- row {row['row_index']}: id={row['id']} name={row['name']}")

    if missing_ids:
        print("Missing ids:")
        for entry in missing_ids:
            print(f"- {entry['id']} ({entry['name']})")

    if duplicates:
        print("Duplicates report:", REPORTS_DIR / f"{base_name}__duplicates.csv")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
